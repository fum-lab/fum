"""Полная сборка PPTX и PDF из единого содержания слайдов."""

from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph


def шрифты(обычный, жирный):
    pdfmetrics.registerFont(TTFont("FUMA", str(обычный)))
    pdfmetrics.registerFont(TTFont("FUMA-Bold", str(жирный)))


def собрать_презентацию(данные, презентация, выход):
    ширина, высота = 960, 540
    книга = Presentation()
    книга.slide_width, книга.slide_height = Inches(13.333333), Inches(7.5)
    книга.core_properties.title = презентация["название"].replace("\n", " ")
    книга.core_properties.author = "FUM"
    книга.core_properties.created = datetime(2026, 9, 23, tzinfo=timezone.utc)
    книга.core_properties.modified = книга.core_properties.created
    файл = Path(выход) / презентация["файл"]
    холст = Canvas(str(файл.with_suffix(".pdf")), pagesize=(ширина, высота), invariant=1)
    холст.setTitle(книга.core_properties.title)
    холст.setAuthor("FUM")
    страницы = []

    def страница(тёмная=False):
        слайд = книга.slides.add_slide(книга.slide_layouts[6])
        цвет = "102436" if тёмная else "F5F3ED"
        слайд.background.fill.solid()
        слайд.background.fill.fore_color.rgb = RGBColor.from_string(цвет)
        холст.setFillColor(HexColor("#" + цвет))
        холст.rect(0, 0, ширина, высота, fill=1, stroke=0)
        return слайд

    def блок(слайд, текст, х, у, ш, в, размер=18, жирный=False, цвет="102436", ссылка=None):
        область = слайд.shapes.add_textbox(Inches(х / 72), Inches(у / 72), Inches(ш / 72), Inches(в / 72))
        рамка = область.text_frame
        рамка.word_wrap = True
        рамка.margin_left = рамка.margin_right = рамка.margin_top = рамка.margin_bottom = 0
        for номер, строка in enumerate(текст.split("\n")):
            абзац = рамка.paragraphs[0] if номер == 0 else рамка.add_paragraph()
            абзац.text = строка
            абзац.font.name = "Arial"
            абзац.font.size = Pt(размер)
            абзац.font.bold = жирный
            абзац.font.color.rgb = RGBColor.from_string(цвет)
            абзац.space_after = Pt(3)
            if ссылка:
                for фрагмент in абзац.runs:
                    фрагмент.hyperlink.address = ссылка
        стиль = ParagraphStyle("FUMA", fontName="FUMA-Bold" if жирный else "FUMA", fontSize=размер,
                              leading=размер * 1.22, textColor=HexColor("#" + цвет))
        содержимое = escape(текст).replace("\n", "<br/>")
        if ссылка:
            содержимое = '<link href="' + escape(ссылка, {'"': '&quot;'}) + '">' + содержимое + '</link>'
        абзац = Paragraph(содержимое, стиль)
        _, фактическая_высота = абзац.wrap(ш, в)
        if фактическая_высота > в + 0.1:
            raise ValueError("Текст не помещается в PDF: " + текст[:100])
        абзац.drawOn(холст, х, высота - у - фактическая_высота)

    def колонтитул(слайд, номер, тёмная=False):
        цвет = "B8C7D0" if тёмная else "52626A"
        блок(слайд, "FUMA  /  " + данные["дата"] + "  /  Предложение для обсуждения", 44, 505, 810, 18, 9, цвет=цвет)
        блок(слайд, str(номер).zfill(2), 891, 502, 30, 20, 12, цвет=цвет)

    слайд = страница(True)
    блок(слайд, "FUMA", 44, 42, 870, 42, 25, True, "C5EE90")
    блок(слайд, презентация["адресат"], 44, 106, 870, 34, 16, цвет="B8C7D0")
    блок(слайд, презентация["название"], 44, 170, 870, 170, 42, True, "FFFFFF")
    блок(слайд, презентация["подзаголовок"], 44, 370, 810, 78, 20, цвет="DDE6EC")
    колонтитул(слайд, 1, True)
    холст.showPage()
    страницы.append(презентация["название"])

    свидетельство = {
        "заголовок": "Текущая зрелость и следующий шаг",
        "тезисы": [
            {"заголовок": "Работающий MVP", "текст": "Связка памяти FUM и Codex с участием человека: постановки, изменения, проверки и сохранение результатов."},
            {"заголовок": "Проверенный участок", "текст": "SwiftPM и Xcode: исходные байты сохраняются до разбора; исполнение оператора и повтор проверены на ограниченном файловом вводе."},
            {"заголовок": "Целевой продукт", "текст": "Голос, аватар, генерация интерфейсов, непрерывная память и широкое управление устройствами ещё требуют реализации и проверки."}
        ],
        "вопрос": "Контрольная точка исходников: 2c0ccc7 · 8 тестов исполнения и реальные CLI-сценарии",
        "источники": []
    }
    for номер, данные_слайда in enumerate([*презентация["слайды"], свидетельство], 2):
        слайд = страница()
        блок(слайд, презентация["адресат"].upper(), 44, 30, 865, 28, 10, цвет="52626A")
        блок(слайд, данные_слайда["заголовок"], 44, 83, 872, 92, 29, True)
        for колонка, тезис in enumerate(данные_слайда["тезисы"]):
            х = 44 + колонка * 299
            блок(слайд, str(колонка + 1).zfill(2), х, 195, 275, 34, 19, True, "648331")
            блок(слайд, тезис["заголовок"], х, 238, 268, 61, 20, True)
            блок(слайд, тезис["текст"], х, 311, 268, 140, 16)
        вопрос = данные_слайда.get("вопрос")
        if вопрос:
            блок(слайд, вопрос, 44, 462, 872, 32, 12, True, "405E20")
        колонтитул(слайд, номер)
        ссылки = [данные["источники"][имя] for имя in данные_слайда.get("источники", [])]
        слайд.notes_slide.notes_text_frame.text = "\n".join(["Статус: предложение; партнёрство и пилот не подтверждены.", *[с["название"] + ": " + с["url"] for с in ссылки]])
        холст.showPage()
        страницы.append(данные_слайда["заголовок"])

    источники = list(dict.fromkeys(имя for с in презентация["слайды"] for имя in с.get("источники", [])))
    слайд = страница(True)
    блок(слайд, "Основания и контакты следующего шага", 44, 54, 870, 95, 29, True, "FFFFFF")
    for номер, имя in enumerate(источники):
        источник = данные["источники"][имя]
        блок(слайд, источник["название"], 44, 165 + номер * 35, 870, 30, 16, цвет="DDE6EC", ссылка=источник["url"])
    блок(слайд, "Контакт со стороны FUM: назначается учредителями.\nЮрлицо, сумма и условия сотрудничества ещё согласуются.", 44, 420, 870, 66, 14, цвет="B8C7D0")
    колонтитул(слайд, len(страницы) + 1, True)
    холст.showPage()
    холст.save()
    книга.save(str(файл.with_suffix(".pptx")))
    страницы.append("Основания")
    текст = ["# " + презентация["название"].replace("\n", " "), "", презентация["адресат"], "", презентация["подзаголовок"]]
    for с in [*презентация["слайды"], свидетельство]:
        текст += ["", "## " + с["заголовок"], ""]
        текст += ["- **" + т["заголовок"] + ".** " + т["текст"] for т in с["тезисы"]]
        if с.get("вопрос"):
            текст += ["", с["вопрос"]]
    текст += ["", "## Источники", ""] + ["- [" + данные["источники"][и]["название"] + "](" + данные["источники"][и]["url"] + ")" for и in источники]
    файл.with_suffix(".md").write_text("\n".join(текст) + "\n", encoding="utf-8")
    return {"название": презентация["адресат"], "файл": презентация["файл"], "страниц": len(страницы)}
