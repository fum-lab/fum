"""Производные DOCX и PDF из открытых Markdown-документов пакета."""

import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


def простой_текст(строка):
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", строка).replace("**", "").replace("`", "")


def собрать_документ(исходник, выход):
    строки = исходник.read_text(encoding="utf-8").splitlines()
    документ = Document()
    документ.styles["Normal"].font.name = "Arial"
    документ.styles["Normal"].font.size = Pt(10)
    стили = {
        "текст": ParagraphStyle("текст", fontName="FUMA", fontSize=10, leading=15, spaceAfter=8),
        "заголовок": ParagraphStyle("заголовок", fontName="FUMA-Bold", fontSize=22, leading=28, spaceAfter=20),
        "раздел": ParagraphStyle("раздел", fontName="FUMA-Bold", fontSize=14, leading=18, spaceBefore=15, spaceAfter=10),
        "ячейка": ParagraphStyle("ячейка", fontName="FUMA", fontSize=8, leading=11),
    }
    поток = []
    номер = 0
    while номер < len(строки):
        строка = строки[номер].strip()
        номер += 1
        if not строка or строка.startswith("<!--"):
            continue
        if строка.startswith("|"):
            таблица = []
            while True:
                ячейки = [ячейка.strip() for ячейка in строка.strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", ячейка) for ячейка in ячейки):
                    таблица.append([простой_текст(ячейка) for ячейка in ячейки])
                if номер >= len(строки) or not строки[номер].strip().startswith("|"):
                    break
                строка = строки[номер].strip()
                номер += 1
            таблица_docx = документ.add_table(rows=1, cols=len(таблица[0]))
            таблица_docx.style = "Light Shading Accent 1"
            for столбец, текст in enumerate(таблица[0]):
                таблица_docx.rows[0].cells[столбец].text = текст
            for ряд in таблица[1:]:
                for ячейка, текст in zip(таблица_docx.add_row().cells, ряд):
                    ячейка.text = текст
            таблица_pdf = Table([[Paragraph(escape(ячейка), стили["ячейка"]) for ячейка in ряд] for ряд in таблица],
                               colWidths=[495 / len(таблица[0])] * len(таблица[0]), repeatRows=1)
            таблица_pdf.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE8D4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BAC5CA")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            поток.extend([таблица_pdf, Spacer(1, 12)])
            continue
        if строка.startswith("# "):
            текст = простой_текст(строка[2:])
            документ.add_heading(текст, 0)
            поток.append(Paragraph(escape(текст), стили["заголовок"]))
        elif строка.startswith("## "):
            текст = простой_текст(строка[3:])
            документ.add_heading(текст, 1)
            поток.append(Paragraph(escape(текст), стили["раздел"]))
        else:
            текст = простой_текст(строка)
            документ.add_paragraph(текст)
            поток.append(Paragraph(escape(текст), стили["текст"]))
    документ.save(str(выход.with_suffix(".docx")))
    пдф = SimpleDocTemplate(str(выход.with_suffix(".pdf")), pagesize=A4, leftMargin=50, rightMargin=50,
                           topMargin=45, bottomMargin=45, title=простой_текст(строки[0]), author="FUM")
    пдф.build(поток)
